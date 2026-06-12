"""Excel I/O and reporting helpers."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from config import OUTPUT_COLUMNS
from scraper import SchoolRecord


def read_locations(path: Path) -> list[str]:
    if not path.exists():
        raise FileNotFoundError(f"Locations file not found: {path}")

    df = pd.read_excel(path)
    if df.empty:
        raise ValueError(f"No locations found in {path}")

    # Accept common column names for the location list.
    for column in ("Location", "Location List", "location", "City", "city"):
        if column in df.columns:
            values = [str(v).strip() for v in df[column].tolist() if str(v).strip() and str(v).lower() != "nan"]
            if values:
                return values

    # Fallback: first column.
    first_col = df.columns[0]
    values = [str(v).strip() for v in df[first_col].tolist() if str(v).strip() and str(v).lower() != "nan"]
    if not values:
        raise ValueError(f"Could not read locations from {path}")
    return values


def ensure_template(path: Path) -> None:
    if path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Schools"
    ws.append(OUTPUT_COLUMNS)
    wb.save(path)


def write_output(records: list[SchoolRecord], template_path: Path, output_path: Path) -> None:
    ensure_template(template_path)
    rows = [record.to_row() for record in records]
    df = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)

    template_wb = load_workbook(template_path)
    sheet_name = template_wb.sheetnames[0]
    template_wb.close()

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)


def write_missing_report(
    path: Path,
    records: list[SchoolRecord],
    location_errors: list[str],
    duplicate_notes: list[str],
) -> None:
    lines: list[str] = [
        "School Scraper - Missing Data Report",
        f"Generated: {datetime.now().isoformat(timespec='seconds')}",
        "",
        f"Total schools collected: {len(records)}",
        "",
    ]

    if location_errors:
        lines.append("=== Location Errors ===")
        lines.extend(f"- {note}" for note in location_errors)
        lines.append("")

    if duplicate_notes:
        lines.append("=== Duplicates Removed ===")
        lines.extend(f"- {note}" for note in duplicate_notes)
        lines.append("")

    missing_rows = [r for r in records if r.missing_fields]
    lines.append(f"=== Schools With Missing Fields ({len(missing_rows)}) ===")
    for record in missing_rows:
        fields = ", ".join(record.missing_fields)
        lines.append(f"- {record.school or record.profile_url}: missing {fields}")

    schools_without_website = [r.school for r in records if not r.website]
    if schools_without_website:
        lines.append("")
        lines.append(f"=== Missing Website ({len(schools_without_website)}) ===")
        lines.extend(f"- {name}" for name in schools_without_website[:100])
        if len(schools_without_website) > 100:
            lines.append(f"... and {len(schools_without_website) - 100} more")

    path.write_text("\n".join(lines), encoding="utf-8")
